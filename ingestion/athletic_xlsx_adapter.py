"""
Phase 4C real adapter: The Athletic season-long projections/rankings xlsx.

This is the source-specific acquisition/normalization code for the one real
specimen captured so far (a subscriber-distributed Excel draft-prep model,
season-long/preseason in scope — not weekly). Per the frozen boundary: this
file is responsible ONLY for reading the workbook and producing
NormalizedObservations. It contains no trust or persistence logic — every
observation it emits still goes through the exact same
ingestion.pipeline.ingest_collector_batch() -> _route_field() path any other
source does. Identity resolution, validation, confidence routing, versioning,
and conflict detection are unchanged and untouched by this file.

Known limitations of THIS adapter, stated plainly rather than glossed over:
- Season-long/preseason scope only (scope='season') — this workbook has no
  per-week breakdown, so it cannot produce weekly observations.
- Raw counting stats only (attempts, yards, TDs, completions, receptions) —
  the workbook's own precomputed FPS/AUC$/HALF/PPR columns are deliberately
  NOT imported, since raw stats are authoritative here and this app always
  derives fantasy points itself (never trusts a source's own point total).
- fumbles_lost is not present anywhere in this workbook for any position —
  those observations are simply never produced; the app already treats an
  absent stat as "missing", not zero.
- Rankings are read from the "Rankings" sheet's first QB/RB/WR/TE column
  group only (positional ranks). The sheet's later duplicate RB/WR/TE groups
  and the separate "OVR & VORP Ranks" sheet (which would give a genuine
  cross-position "overall" rank) are NOT parsed in this pass — out of scope
  until there's a clear, confirmed reason to interpret them.
"""
from dataclasses import replace
from typing import List

from ingestion.collector import SourceAdapter, NormalizedObservation, AdapterUnavailable

# sheet column header -> our internal stat_name. Deliberately excludes
# columns that are targets/attempts-we-don't-track (RUAT, TGT) or precomputed
# point totals (FPS, HALF, PPR, Custom, AUC$) — raw stats only.
_QB_STAT_MAP = {"PAYD": "pass_yards", "PATD": "pass_tds", "INT": "interceptions",
                "RUYD": "rush_yards", "RUTD": "rush_tds", "CMP": "pass_completions", "PATT": "pass_attempts"}
_RB_STAT_MAP = {"RUYD": "rush_yards", "RUTD": "rush_tds", "REC": "receptions",
                "RCYD": "receiving_yards", "RCTD": "receiving_tds"}
_WR_STAT_MAP = {"RUYD": "rush_yards", "RUTD": "rush_tds", "REC": "receptions",
                "RCYD": "receiving_yards", "RCTD": "receiving_tds"}
_TE_STAT_MAP = {"REC": "receptions", "RCYD": "receiving_yards", "RCTD": "receiving_tds"}

_POSITION_SHEETS = {"QB": _QB_STAT_MAP, "RB": _RB_STAT_MAP, "WR": _WR_STAT_MAP, "TE": _TE_STAT_MAP}

# Confidence for this adapter: real structured spreadsheet cells, not OCR/vision
# uncertainty — but still a personal analyst's projection (not a verified official
# stat), and this adapter has no per-field signal of its own. A flat, moderately
# high value reflects "the read itself is reliable" without claiming the
# PROJECTION is certain to be correct — that's what auto-accept + Review Queue
# + versioning is for downstream, not this adapter's job to judge.
_PROJECTION_CONFIDENCE = 0.9
_RANKING_CONFIDENCE = 0.9

_VALID_RANKING_TYPES = {"overall", "QB", "RB", "WR", "TE", "FLEX"}


class AthleticSeasonXlsxAdapter(SourceAdapter):
    """Reads QB/RB/WR/TE sheets for season-long raw stat projections, and the
    Rankings sheet for season-long positional ranks. fetch()'s week_number
    argument is accepted for interface compatibility but ignored — every
    observation this adapter produces is scope='season'."""

    def __init__(self, source_name: str, file_path: str):
        self.source_name = source_name
        self.file_path = file_path

    def fetch(self, season_id: int, week_number: int) -> List[NormalizedObservation]:
        try:
            import openpyxl
        except ImportError as e:
            raise AdapterUnavailable("openpyxl is not installed (pip install openpyxl)", transient=False) from e

        try:
            wb = openpyxl.load_workbook(self.file_path, data_only=True)
        except FileNotFoundError as e:
            raise AdapterUnavailable(f"Source file not found: {self.file_path}", transient=True) from e
        except Exception as e:
            raise AdapterUnavailable(f"Could not open workbook: {e}", transient=False) from e

        observations: List[NormalizedObservation] = []
        observations.extend(self._extract_projections(wb))
        observations.extend(self._extract_rankings(wb))
        return observations

    def _extract_projections(self, wb) -> List[NormalizedObservation]:
        observations = []
        for position, stat_map in _POSITION_SHEETS.items():
            if position not in wb.sheetnames:
                continue
            ws = wb[position]
            header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
            col_index = {name: i for i, name in enumerate(header)}
            if "Player" not in col_index:
                continue  # sheet doesn't match the expected shape — skip rather than guess
            for row in ws.iter_rows(min_row=2, values_only=True):
                player_name = row[col_index["Player"]]
                if not player_name:
                    continue
                for header_key, stat_name in stat_map.items():
                    if header_key not in col_index:
                        continue
                    value = row[col_index[header_key]]
                    if value is None:
                        continue
                    observations.append(NormalizedObservation(
                        player_name_raw=str(player_name),
                        position_hint=position,
                        data_type="projection",
                        stat_name=stat_name,
                        value=round(float(value), 2),
                        confidence=_PROJECTION_CONFIDENCE,
                        scope="season",
                        team_hint=row[col_index["TM"]] if "TM" in col_index else None,
                    ))
        return observations

    def _extract_rankings(self, wb) -> List[NormalizedObservation]:
        observations = []
        if "Rankings" not in wb.sheetnames:
            return observations
        ws = wb["Rankings"]
        rows = list(ws.iter_rows(min_row=1, max_row=2, values_only=True))
        if len(rows) < 2:
            return observations
        position_row, field_row = rows[0], rows[1]

        # Find the FIRST column group for each of QB/RB/WR/TE (see module docstring —
        # later duplicate groups and true cross-position "overall" rank are out of
        # scope for this pass).
        group_starts = {}
        for i, label in enumerate(position_row):
            if label in _POSITION_SHEETS and label not in group_starts:
                group_starts[label] = i

        for ranking_type, start_col in group_starts.items():
            # within a group: [RK is col 0 globally] Name, Team, Position, Player ID
            try:
                name_col = start_col + field_row[start_col:].index("Name")
                team_col = start_col + field_row[start_col:].index("Team")
            except ValueError:
                continue  # group doesn't have the expected shape — skip rather than guess
            for row in ws.iter_rows(min_row=3, values_only=True):
                rank_value = row[0]
                player_name = row[name_col]
                if rank_value is None or not player_name:
                    continue
                observations.append(NormalizedObservation(
                    player_name_raw=str(player_name),
                    position_hint=ranking_type if ranking_type in ("QB", "RB", "WR", "TE") else None,
                    data_type="ranking",
                    stat_name="ranking",
                    value=float(rank_value),
                    confidence=_RANKING_CONFIDENCE,
                    scope="season",
                    ranking_type=ranking_type,
                    team_hint=row[team_col] if team_col < len(row) else None,
                ))
        return observations
